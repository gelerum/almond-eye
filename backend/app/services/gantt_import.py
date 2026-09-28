"""Calendar plan import: XLSX Gantt in the organiser's template layout or CSV with explicit dates.

XLSX layout («График производства работ», data/plans/*.xlsx, scripts/make_demo_plans.py):
title rows, a row with years and a row with month names, first column = stage name, filled month cells =
planned interval. Optional columns «Начало» / «Окончание» (exact dates, preferred) and «Ключ этапа»
(stage_key); without the key a row is mapped by plan_keywords of stage_signatures.yaml.
CSV: a header with a name column (work_name / name / Наименование) and start / end columns
(start_date / start / Начало, end_date / end / Окончание), optionally stage_key; the timeline template
(`stage_id;work_name;zone_id;start_date;end_date;rule_profile`) is accepted as is.
Nothing is stored here: the result is a preview for explicit confirmation.
"""
import calendar
import csv
import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime

from openpyxl import load_workbook

from backend.app.domain.stage_inference import load_signatures, match_plan_row
from backend.app.services.import_schedule import MAX_UPLOAD

METHOD = "gantt-plan-v1"
MONTHS = ["январ", "феврал", "март", "апрел", "ма", "июн", "июл", "август", "сентябр", "октябр", "ноябр", "декабр"]
MONTH_SHORT = ["янв", "фев", "мар", "апр", "май", "июн", "июл", "авг", "сен", "окт", "ноя", "дек"]
FILL_MARKS = {"x", "х", "■", "█", "+", "1", "v", "✓"}


def month_index(value) -> int | None:
    if isinstance(value, (datetime, date)):
        return value.month
    text = str(value or "").strip().lower().replace("ё", "е")
    if not text or len(text) > 20:
        return None
    word = re.split(r"[\s.\-/]", text)[0]
    for i, (full, abbr) in enumerate(zip(MONTHS, MONTH_SHORT)):
        if word == "май" or word == "мая":
            return 5
        if i != 4 and (word.startswith(full) or word == abbr):
            return i + 1
    return None


def parse_date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for pattern in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text[:10], pattern).date()
        except ValueError:
            continue
    raise ValueError(f"Не удалось прочитать дату «{text}»: используйте ДД.ММ.ГГГГ или ГГГГ-ММ-ДД")


def issue(severity, code, message, row=None):
    return {"severity": severity, "code": code, "row": row, "message": message}


def month_end(year: int, month: int) -> date:
    return date(year, month, calendar.monthrange(year, month)[1])


def _stage_key(raw_key, name, sig, row, issues):
    key = str(raw_key or "").strip()
    if key:
        if key in sig["stages"]:
            return key, "column"
        issues.append(issue("warning", "unknown_stage_key", f"Неизвестный ключ этапа «{key}»; сопоставляем по названию", row))
    key = match_plan_row(name, sig)
    if key is None:
        issues.append(issue("warning", "unmapped_stage", f"«{name}» не сопоставлена с этапом методики: "
                                                         "строка сохранится, но по снимкам не оценивается", row))
    return key, "keywords" if key else None


def _is_filled(cell) -> bool:
    value = cell.value
    if value not in (None, "") and str(value).strip().lower() in FILL_MARKS:
        return True
    fill = cell.fill
    if not fill or fill.fill_type in (None, "none"):
        return False
    color = fill.fgColor
    if color is None:
        return False
    if color.type == "rgb":
        return str(color.rgb).upper() not in ("00000000", "FFFFFFFF", "FFFFFF")
    return color.type in ("theme", "indexed") and not (color.type == "indexed" and color.indexed in (64, 65))


def _xlsx_rows(content: bytes):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(f.file_size for f in archive.infolist()) > 50 * 1024 * 1024:
            raise ValueError("Распакованный XLSX превышает 50 МБ")
    book = load_workbook(io.BytesIO(content), data_only=True)
    ws = book.worksheets[0]
    if ws.max_row > 2000 or ws.max_column > 400:
        raise ValueError("Лимит плана: 2000 строк и 400 столбцов")
    return ws


def preview_xlsx(content: bytes, sig: dict) -> dict:
    ws = _xlsx_rows(content)
    grid = [[c for c in row] for row in ws.iter_rows()]
    issues = []
    # Строка месяцев: не меньше трёх ячеек с названием месяца (или датой).
    month_row = next((r for r, row in enumerate(grid) if sum(month_index(c.value) is not None
                                                            and not isinstance(c.value, (int, float)) for c in row) >= 3), None)
    titles = [str(c.value).strip() for row in grid[:max(0, (month_row or 3) - 1)] for c in row
              if isinstance(c.value, str) and c.value.strip()]
    header_rows = grid[:(month_row + 1) if month_row is not None else 6]
    columns = {}
    for row in header_rows:
        for c in row:
            text = str(c.value or "").strip().lower()
            for field, words in (("start", ("начало",)), ("end", ("окончание", "конец")),
                                 ("key", ("ключ этапа", "stage_key")), ("name", ("наименование", "этап", "работ"))):
                if field not in columns and any(text.startswith(w) or text == w for w in words):
                    columns[field] = c.column - 1
    name_col = columns.get("name", 0)
    month_cols = {}
    if month_row is not None:
        year = None
        years_row = grid[month_row - 1] if month_row > 0 else []
        year_by_col = {}
        for c in years_row:
            found = re.search(r"(20\d\d)", str(c.value or ""))
            if found:
                year = int(found.group(1))
            if year:
                year_by_col[c.column - 1] = year
        previous = None
        for c in grid[month_row]:
            m = month_index(c.value)
            if m is None or isinstance(c.value, (int, float)) or c.column - 1 in columns.values():
                continue
            y = c.value.year if isinstance(c.value, (datetime, date)) else year_by_col.get(c.column - 1)
            if y is None:
                found = re.search(r"(20\d\d)", str(c.value))
                y = int(found.group(1)) if found else (previous[0] + (m < previous[1]) if previous else None)
            if y is None:
                issues.append(issue("warning", "month_without_year", f"Столбец {c.column_letter}: месяц без года"))
                continue
            month_cols[c.column - 1] = (y, m)
            previous = (y, m)
    if month_row is None and not {"start", "end"} <= columns.keys():
        raise ValueError("Не найдены ни строка месяцев, ни столбцы «Начало»/«Окончание»")
    first_data = (month_row + 1) if month_row is not None else max(r for r, row in enumerate(header_rows)) + 1
    stages = []
    for r in range(first_data, len(grid)):
        row = grid[r]
        excel_row = r + 1
        name = row[name_col].value if name_col < len(row) else None
        if name is None or not str(name).strip():
            continue
        name = " ".join(str(name).split())
        filled = sorted(month_cols[i] for i in month_cols if i < len(row) and _is_filled(row[i]))
        try:
            start = parse_date(row[columns["start"]].value) if "start" in columns else None
            end = parse_date(row[columns["end"]].value) if "end" in columns else None
        except ValueError as exc:
            issues.append(issue("error", "invalid_date", str(exc), excel_row))
            continue
        source = "exact"
        if bool(start) != bool(end):
            issues.append(issue("error", "one_date", "Укажите и начало, и окончание", excel_row))
            continue
        if not start:
            if not filled:
                issues.append(issue("warning", "no_dates", f"«{name}»: нет дат и залитых месяцев — строка пропущена", excel_row))
                continue
            start, end, source = date(*filled[0], 1), month_end(*filled[-1]), "months"
        elif filled and (filled[0] != (start.year, start.month) or filled[-1] != (end.year, end.month)):
            issues.append(issue("warning", "fill_mismatch", f"«{name}»: залитые месяцы не совпадают с датами "
                                                            "«Начало»/«Окончание»; используем точные даты", excel_row))
        if end < start:
            issues.append(issue("error", "end_before_start", f"«{name}»: окончание раньше начала", excel_row))
            continue
        key_value = row[columns["key"]].value if "key" in columns and columns["key"] < len(row) else None
        key, mapped = _stage_key(key_value, name, sig, excel_row, issues)
        stages.append({"row": excel_row, "name": name, "stage_key": key, "mapped_by": mapped,
                       "start": start.isoformat(), "end": end.isoformat(), "dates_from": source})
    title = next((t for t in titles if "график" in t.lower()), titles[0] if titles else None)
    object_name = next((t for t in titles if t != title), None)
    return {"format": "xlsx", "title": title, "object_name": object_name, "stages": stages, "issues": issues}


def preview_csv(content: bytes, sig: dict) -> dict:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError("Сохраните CSV в кодировке UTF-8") from exc
    lines = text.splitlines()
    if not lines:
        raise ValueError("CSV пуст")
    delimiter = ";" if ";" in lines[0] else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    fields = {(f or "").strip().lower(): f for f in reader.fieldnames or []}

    def pick(*names):
        return next((fields[n] for n in names if n in fields), None)
    name_f, start_f = pick("work_name", "name", "наименование", "этап"), pick("start_date", "start", "начало")
    end_f, key_f = pick("end_date", "end", "окончание"), pick("stage_key", "ключ этапа", "stage_id")
    if not (name_f and start_f and end_f):
        raise ValueError("Нужны столбцы названия и сроков: work_name;start_date;end_date (или Наименование;Начало;Окончание)")
    stages, issues = [], []
    for row in reader:
        number = reader.line_num
        name = " ".join(str(row.get(name_f) or "").split())
        if not name:
            continue
        try:
            start, end = parse_date(row.get(start_f)), parse_date(row.get(end_f))
            if not (start and end):
                raise ValueError("Обе даты обязательны")
            if end < start:
                raise ValueError("Окончание раньше начала")
        except ValueError as exc:
            issues.append(issue("error", "invalid_date", f"«{name}»: {exc}", number))
            continue
        key, mapped = _stage_key(row.get(key_f) if key_f else None, name, sig, number, issues)
        stages.append({"row": number, "name": name, "stage_key": key, "mapped_by": mapped,
                       "start": start.isoformat(), "end": end.isoformat(), "dates_from": "exact"})
    return {"format": "csv", "title": None, "object_name": None, "stages": stages, "issues": issues}


def preview_plan(content: bytes, filename: str, signatures: dict | None = None) -> dict:
    """Parse a plan file into a schedule payload with issues; raises ValueError on unreadable input."""
    if not content or len(content) > MAX_UPLOAD:
        raise ValueError("Файл пуст или превышает 10 МБ")
    sig = signatures or load_signatures()
    lower = (filename or "").lower()
    try:
        if lower.endswith(".xlsx"):
            result = preview_xlsx(content, sig)
        elif lower.endswith(".csv"):
            result = preview_csv(content, sig)
        else:
            raise ValueError("Поддерживаются календарные планы XLSX и CSV")
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError("Не удалось прочитать файл плана") from exc
    stages, issues = result["stages"], result["issues"]
    if not stages:
        issues.append(issue("error", "empty", "В плане не найдено ни одного этапа со сроками"))
    elif not any(s["stage_key"] and sig["stages"][s["stage_key"]].get("observable", True) for s in stages):
        issues.append(issue("warning", "nothing_observable", "Ни один этап плана не распознаётся по технике на снимках"))
    for key in {s["stage_key"] for s in stages if s["stage_key"]}:
        rows = [s for s in stages if s["stage_key"] == key]
        if len(rows) > 1 and sig["stages"][key].get("observable", True):
            issues.append(issue("warning", "duplicate_stage", f"Этап «{sig['stages'][key]['name']}» встречается в "
                                f"{len(rows)} строках: снимки будут засчитаны каждой из них", rows[1]["row"]))
    payload = {"method_version": METHOD, "signatures_version": sig.get("version"), "filename": filename,
               "file_sha256": hashlib.sha256(content).hexdigest(), **result,
               "start": min((s["start"] for s in stages), default=None),
               "end": max((s["end"] for s in stages), default=None),
               # Совместимость с формой графика (schedule): works с датами, календарь привязан.
               "calendar_bound": True, "parameters": {"weekdays": list(range(7)), "holidays": []},
               "works": [{"stage_id": s["stage_key"] or f"row:{s['row']}", "stage_key": s["stage_key"], "name": s["name"],
                          "source_row": s["row"], "start_date": s["start"], "end_date": s["end"], "zone_id": None,
                          "provenance": "uploaded", "rule": None} for s in stages]}
    payload["fingerprint"] = hashlib.sha256(json.dumps(payload, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    return payload


def plan_rows(plan: dict) -> list[dict]:
    """Stages of a stored plan document in the shape progress.assess expects."""
    return [{"stage_key": s["stage_key"], "name": s["name"], "start": s["start"], "end": s["end"], "row": s["row"]}
            for s in plan["stages"]]




def render_gantt_xlsx(object_name: str, rows: list[dict], title: str = "График производства работ",
                      hide_key: bool = True) -> bytes:
    """Rows: {name, start, end, stage_key?} → XLSX in the organiser's template layout + exact dates and stage key."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    starts = [date.fromisoformat(str(r["start"])) for r in rows]
    ends = [date.fromisoformat(str(r["end"])) for r in rows]
    months, (y, m) = [], (min(starts).year, min(starts).month)
    while (y, m) <= (max(ends).year, max(ends).month):
        months.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    wb = Workbook()
    ws = wb.active
    ws.title = "График"
    fixed = ["№", "Наименование работ", "Начало", "Окончание", "Ключ этапа"]
    offset, last_col = len(fixed) + 1, len(fixed) + len(months)
    thin = Side(style="thin", color="999999")
    border, bold = Border(left=thin, right=thin, top=thin, bottom=thin), Font(bold=True)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(row=2, column=1, value=f"Объект: {object_name}").font = Font(italic=True)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    for col, text in enumerate(fixed, 1):
        cell = ws.cell(row=4, column=col, value=text)
        cell.font, cell.alignment, cell.border = bold, center, border
        ws.merge_cells(start_row=4, start_column=col, end_row=5, end_column=col)
    col = offset
    for year in sorted({y for y, _ in months}):  # строка лет: объединённые ячейки над месяцами года
        span = sum(1 for y, _ in months if y == year)
        cell = ws.cell(row=4, column=col, value=year)
        cell.font, cell.alignment, cell.border = bold, center, border
        if span > 1:
            ws.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + span - 1)
        col += span
    for i, (y, m) in enumerate(months):
        cell = ws.cell(row=5, column=offset + i, value=MONTH_SHORT[m - 1])
        cell.alignment, cell.border = center, border
        ws.column_dimensions[get_column_letter(offset + i)].width = 5
    fill = PatternFill("solid", fgColor="4F81BD")
    for r, (row, start, end) in enumerate(zip(rows, starts, ends), 6):
        for col, value in enumerate([r - 5, row["name"], start, end, row.get("stage_key") or ""], 1):
            cell = ws.cell(row=r, column=col, value=value)
            cell.border = border
            if isinstance(value, date):
                cell.number_format = "DD.MM.YYYY"
        for i, (y, m) in enumerate(months):
            cell = ws.cell(row=r, column=offset + i)
            cell.border = border
            if (start.year, start.month) <= (y, m) <= (end.year, end.month):
                cell.fill = fill
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 48
    ws.column_dimensions["C"].width = ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 20
    ws.column_dimensions["E"].hidden = hide_key
    ws.freeze_panes = ws.cell(row=6, column=offset)
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()
